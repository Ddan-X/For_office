import os
from docx import Document

fl =[]
jidu=0
tianzhu=0
fojiao=0
daojiao=0
qita=0

#获取文件
def get_file(filepath):
    for i in os.listdir(filepath):
        path = os.path.join(filepath, i)
        if os.path.isdir(path):
            get_file(path)
        if path.endswith(".docx"):
            #print(path)
            fl.append(path)
get_file('F:\python_DanL\ZoeOffice\登记表') #存放文件的文件夹路径，根据不同进行修改
#print(fl)
for f in fl:
    # print(f)
    document = Document(f) #docx文件
    tables = document.tables
    table = tables[0]  # 获取文件中的第一个表格

    for i in range(1, len(table.rows)):  # 从表格第二行开始循环读取表格数据
        result = table.cell(i, 0).text + "" + table.cell(i, 1).text + table.cell(i, 2).text + table.cell(i, 3).text
    # print(table.cell(i,0).text)
        #print(table.cell(i,1).text)
    # cell(i,0)表示第(i+1)行第1列数据，以此类推
    # print(result)
        for info in table.cell(i, 1).text:
            #print(info)
            if info == '基':
                jidu = jidu + 1

            elif info == '天':
                tianzhu = tianzhu + 1

            elif info == '佛':
                fojiao = fojiao + 1

            elif info == '道':
                daojiao = daojiao + 1

            elif info == '其他':
                qita = qita + 1

print('基督：' ,jidu)
print('天主：' , tianzhu)
print('佛：' ,fojiao)
print('道：' ,daojiao)
print('其他：' , qita)